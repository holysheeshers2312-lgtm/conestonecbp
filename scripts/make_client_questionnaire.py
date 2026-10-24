"""Generate ConeStone CBP client intake questionnaire as a Word document."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parents[2] / "ConeStone_CBP_Website_Client_Questionnaire.docx"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_styled(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True, color=RGBColor(0x8B, 0x1E, 0x1E))
    return p


def add_body(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_note(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=10, color=RGBColor(0x5C, 0x4A, 0x3D))
    p.paragraph_format.space_after = Pt(8)
    return p


def add_question(doc, number, question, hint=None, choices=None, lines=1):
    p = doc.add_paragraph()
    run = p.add_run(f"{number}. {question}")
    set_run_font(run, size=11, bold=True)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)

    if hint:
        h = doc.add_paragraph()
        hr = h.add_run(hint)
        set_run_font(hr, size=10, color=RGBColor(0x5C, 0x4A, 0x3D))
        h.paragraph_format.space_after = Pt(2)

    if choices:
        c = doc.add_paragraph()
        cr = c.add_run(choices)
        set_run_font(cr, size=11)
        c.paragraph_format.space_after = Pt(2)

    for _ in range(lines):
        blank = doc.add_paragraph()
        br = blank.add_run("________________________________________________")
        set_run_font(br, size=11, color=RGBColor(0x99, 0x99, 0x99))
        blank.paragraph_format.space_after = Pt(2)


def add_checkbox_row(doc, label):
    p = doc.add_paragraph()
    run = p.add_run(f"☐  {label}")
    set_run_font(run, size=11)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.15)


def add_table_checklist(doc, rows):
    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.style = "Table Grid"
    headers = ["What we need", "Do you have it?", "Notes / link"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run_font(run, size=10, bold=True)

    for r_i, (item, _) in enumerate(rows, start=1):
        table.rows[r_i].cells[0].text = item
        table.rows[r_i].cells[1].text = "☐ Yes   ☐ No   ☐ Later"
        table.rows[r_i].cells[2].text = ""
        for cell in table.rows[r_i].cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, size=10)

    doc.add_paragraph()


def shade_cell(cell, hex_color):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    shading.set(qn("w:val"), "clear")
    cell._tePr = cell._tc.get_or_add_tcPr()
    cell._tc.get_or_add_tcPr().append(shading)


def main():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("ConeStone CBP  ·  Amma Kitchen")
    set_run_font(r, size=20, bold=True, color=RGBColor(0x8B, 0x1E, 0x1E))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Website information form")
    set_run_font(sr, size=14, bold=True)

    add_note(
        doc,
        "Please fill this form so we can finish your website. "
        "Tick the boxes, circle your answers, or write in the blanks. "
        "If you are not sure about something, write “Not sure” or “Later”.",
    )

    add_body(doc, "Your name: ________________________________", size=11)
    add_body(doc, "Role / company: ___________________________", size=11)
    add_body(doc, "WhatsApp / phone: _________________________", size=11)
    add_body(doc, "Email: ____________________________________", size=11)
    add_body(doc, "Date: _____________________________________", size=11)

    # Section A
    add_heading_styled(doc, "A. Logo and brand look", 1)
    add_note(doc, "This helps the website look like your restaurant.")

    add_question(
        doc,
        1,
        "Do you have a logo for ConeStone CBP?",
        choices="☐ Yes, I will send it    ☐ No, use text only for now    ☐ Not sure",
    )
    add_question(
        doc,
        2,
        "Do you have a separate logo for Amma Kitchen?",
        choices="☐ Yes    ☐ No, use ConeStone logo only    ☐ Not sure",
    )
    add_question(
        doc,
        3,
        "If yes, please send the logo file(s).",
        hint="Best formats: SVG or PNG with clear / transparent background. Also send a high-quality version if you have one.",
        lines=2,
    )
    add_question(
        doc,
        4,
        "Do you have special brand colours we must use?",
        choices="☐ Yes (write colours or attach brand guide below)    ☐ No, current website colours are fine",
        lines=2,
    )
    add_question(
        doc,
        5,
        "How should the names appear on the website?",
        choices="☐ ConeStone CBP only    ☐ Amma Kitchen only    ☐ Both names together (recommended)",
    )
    add_question(
        doc,
        6,
        "Keep the line “Traditional Cooking” under the name?",
        choices="☐ Yes, keep it    ☐ No, change it to: _______________________________",
    )

    # Section B
    add_heading_styled(doc, "B. Photos (very important)", 1)
    add_note(
        doc,
        "Real photos make the website feel trustworthy. "
        "Phone photos are fine if they are bright and clear. "
        "Please share via Google Drive, WhatsApp, or email.",
    )

    add_question(
        doc,
        7,
        "Can you share photos of the outside of the restaurant?",
        choices="☐ Yes    ☐ No    ☐ Later",
    )
    add_question(
        doc,
        8,
        "Can you share photos of the inside (seating, counter, bar)?",
        choices="☐ Yes    ☐ No    ☐ Later",
    )
    add_question(
        doc,
        9,
        "Can you share food photos (tiffins, chaat, snacks, special dishes)?",
        choices="☐ Yes    ☐ No    ☐ Later",
    )
    add_question(
        doc,
        10,
        "Can you share drink / coffee photos?",
        choices="☐ Yes    ☐ No    ☐ Later",
    )
    add_question(
        doc,
        11,
        "Any photos we should NOT use?",
        lines=2,
    )
    add_question(
        doc,
        12,
        "Who took the photos, and can we use them on the website?",
        hint="Example: Our staff / a photographer. Yes, you can use them online.",
        lines=2,
    )
    add_question(
        doc,
        13,
        "Where will you send the photos?",
        choices="☐ Google Drive link    ☐ WhatsApp    ☐ Email    ☐ Other: _______________",
        lines=2,
    )

    # Section C
    add_heading_styled(doc, "C. Address, phone and email", 1)
    add_note(doc, "Please check these carefully. Guests will use them to find and contact you.")

    add_question(
        doc,
        14,
        "Is this address correct?",
        hint="6 Changi Business Park Ave 1, #01-23 ESR Park @Changi, Singapore 486017",
        choices="☐ Yes, correct    ☐ No, please change to:",
        lines=2,
    )
    add_question(
        doc,
        15,
        "Public email for the website",
        hint="Current draft on the site: contact@conestonecbp.com",
        choices="☐ Keep this email    ☐ Use a different email:",
        lines=1,
    )
    add_question(
        doc,
        16,
        "Phone numbers on the website",
        hint="Currently listed: Madhav · 8172 5039 and Soujanya · 9336 5499",
        choices="☐ Keep both    ☐ Change / add numbers below",
        lines=2,
    )
    add_question(
        doc,
        17,
        "Which number should be the main contact number?",
        lines=1,
    )
    add_question(
        doc,
        18,
        "Do you want WhatsApp linked on the website?",
        choices="☐ Yes, number: _______________    ☐ No",
    )
    add_question(
        doc,
        19,
        "Exact Google / review name for the restaurant",
        hint="Currently noted as “Cone Stone”. Write the exact name guests should search.",
        lines=1,
    )

    # Section D
    add_heading_styled(doc, "D. Opening hours", 1)
    add_note(
        doc,
        "Hours are not printed in your menu booklet yet. "
        "Please fill this so we can show correct timings on the website.",
    )

    days = [
        "Monday",
        "Tuesday",
        "Wednesday",
        "Thursday",
        "Friday",
        "Saturday",
        "Sunday",
    ]
    table = doc.add_table(rows=1 + len(days), cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(["Day", "Open – Close", "Closed?"]):
        cell = table.rows[0].cells[i]
        cell.text = h
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run_font(run, size=10, bold=True)
    for i, day in enumerate(days, start=1):
        table.rows[i].cells[0].text = day
        table.rows[i].cells[1].text = "____ : ____  to  ____ : ____"
        table.rows[i].cells[2].text = "☐ Closed"
        for cell in table.rows[i].cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, size=10)

    doc.add_paragraph()
    add_question(
        doc,
        20,
        "Any special notes? (public holidays, last order time, kitchen closes earlier, etc.)",
        lines=2,
    )
    add_question(
        doc,
        21,
        "Do you take walk-ins, reservations, or both?",
        choices="☐ Walk-ins only    ☐ Reservations preferred    ☐ Both",
    )

    # Section E
    add_heading_styled(doc, "E. Menu", 1)
    add_note(doc, "Your current menu PDF is already on the website. Tell us if anything changed.")

    add_question(
        doc,
        22,
        "Is the latest menu PDF still correct?",
        choices="☐ Yes    ☐ No — I will send a new PDF",
    )
    add_question(
        doc,
        23,
        "Any dishes to highlight, hide, or mark as often unavailable?",
        lines=2,
    )
    add_question(
        doc,
        24,
        "Should the site say anything about GST / service charge?",
        choices="☐ Prices as shown    ☐ Add this note: _______________________________",
    )
    add_question(
        doc,
        25,
        "Are catering prices still correct? ($24 non-veg / $20 veg per head)",
        choices="☐ Yes    ☐ No — new prices: _______________",
    )

    # Section F
    add_heading_styled(doc, "F. How the website should work", 1)

    add_question(
        doc,
        26,
        "How should guests book a table?",
        choices="☐ Website form (sends email)    ☐ WhatsApp    ☐ Phone call only    ☐ Form + phone",
    )
    add_question(
        doc,
        27,
        "Who should receive booking / enquiry emails?",
        lines=1,
    )
    add_question(
        doc,
        28,
        "Feel of the website",
        choices="☐ Warm family kitchen    ☐ Modern bar & dining    ☐ Both (current direction)",
    )
    add_question(
        doc,
        29,
        "Any website you like (for style inspiration)?",
        lines=2,
    )
    add_question(
        doc,
        30,
        "Anything that must NOT appear on the website?",
        lines=2,
    )

    # Section G
    add_heading_styled(doc, "G. Domain, social media and launch", 1)

    add_question(
        doc,
        31,
        "Do you own the domain conestonecbp.com?",
        choices="☐ Yes    ☐ No    ☐ Not sure",
    )
    add_question(
        doc,
        32,
        "Social media links to show on the website (write full links)",
        hint="Instagram / Facebook / TikTok / others. Leave blank if none.",
        lines=3,
    )
    add_question(
        doc,
        33,
        "Preferred website launch date",
        lines=1,
    )
    add_question(
        doc,
        34,
        "Who will give final approval before we publish?",
        hint="Name + WhatsApp / email",
        lines=2,
    )

    # Section H
    add_heading_styled(doc, "H. Quick checklist — please tick", 1)
    add_note(doc, "This is the short list of what we need most.")

    add_table_checklist(
        doc,
        [
            ("Logo file", ""),
            ("Outside photos", ""),
            ("Inside photos", ""),
            ("Food photos", ""),
            ("Latest menu PDF", ""),
            ("Opening hours", ""),
            ("Final public email", ""),
            ("Social media links", ""),
        ],
    )

    # Section I
    add_heading_styled(doc, "I. Quick choices before launch", 1)
    add_question(
        doc,
        35,
        "Can we launch first with temporary photos, then replace when yours are ready?",
        choices="☐ Yes    ☐ No, wait for our photos",
    )
    add_question(
        doc,
        36,
        "Show opening hours only after you confirm them?",
        choices="☐ Yes    ☐ Hours above are final — show them",
    )
    add_question(
        doc,
        37,
        "Main button on the website should be:",
        choices="☐ Contact us    ☐ Reserve a table    ☐ Both",
    )

    add_heading_styled(doc, "Thank you", 1)
    add_body(
        doc,
        "Please return this form with your files / links. "
        "If anything is unclear, message us and we will help.",
    )
    add_body(doc, "Minimum needed to finish well: logo (or OK to use text), photos, email, hours, latest menu, and one approval contact.", bold=True)

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
